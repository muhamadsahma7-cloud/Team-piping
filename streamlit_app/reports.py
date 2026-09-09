"""Spool classification + Excel exports, ported from the desktop scripts
classify_spools.py and export_master.py. All functions take a DataFrame
(the full `spools` table) so they are database-agnostic.
"""

from __future__ import annotations

import io
from datetime import datetime
from zoneinfo import ZoneInfo

MYT = ZoneInfo("Asia/Kuala_Lumpur")

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------
# classification  (mirrors classify_spools.py)
# ----------------------------------------------------------------------
_SS_RELEASE = {"SS304", "SS_304", "SS316", "SS_316"}
_SS_PAINT_SKIP = {"SS", "SS304", "SS316"}


def classify(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for c in ("line_no", "iso_dwg_no", "dwg_spool_no", "iso_run_no"):
        df[c] = df[c].fillna("").astype(str).str.strip()
    df["spool_key"] = (
        df["line_no"] + "_" + df["iso_dwg_no"] + "_"
        + df["dwg_spool_no"] + "_" + df["iso_run_no"]
    )

    date_cols = [
        "fitup_date", "welding_date", "pwht_date", "fitup_inspection_date",
        "welding_inspection_date", "irn_date", "delivery_date", "site_delivery_date",
    ]
    for c in date_cols:
        if c in df.columns:
            df[c] = df[c].fillna("").astype(str).str.strip()

    status_map: dict[str, str] = {}
    for key, sdf in df.groupby("spool_key"):
        is_straight = sdf["dwg_spool_no"].str.startswith("SP-SPL").all()
        if not is_straight:
            sdf = sdf[sdf["shop_field"] == "S"]
            if sdf.empty:
                continue

        material_group = str(sdf["material_group"].fillna("").iloc[0]).upper()
        pwht_flag = str(sdf["pwht"].fillna("").iloc[0]).upper()
        site_any = sdf["site_delivery_date"].ne("").any()
        paint_any = sdf["delivery_date"].ne("").any()

        fitup_filled = sdf["fitup_date"].ne("").all()
        weld_filled = sdf["welding_date"].ne("").all()
        pwht_filled = sdf["pwht_date"].ne("").all()
        fitup_insp_filled = sdf["fitup_inspection_date"].ne("").all()
        weld_insp_filled = sdf["welding_inspection_date"].ne("").all()
        irn_filled = sdf["irn_date"].ne("").all()

        if is_straight:
            if site_any:
                status = "Sent to Site"
            elif paint_any:
                status = "Sent to Painting"
            else:
                status = "Ready to Release-Straight Pipe"
        elif site_any:
            status = "Sent to Site"
        elif paint_any:
            status = "Sent to Site" if material_group in _SS_PAINT_SKIP else "Sent to Painting"
        elif material_group in _SS_RELEASE and fitup_insp_filled and weld_insp_filled and irn_filled:
            status = "Ready to Release"
        elif pwht_flag == "YES" and fitup_filled and weld_filled and not pwht_filled:
            status = "Ready for PWHT"
        elif (pwht_flag == "YES" and fitup_filled and weld_filled and pwht_filled) or \
             (pwht_flag != "YES" and fitup_filled and weld_filled):
            status = "Ready to Release" if irn_filled else "All done-Awaiting IRN"
        elif fitup_insp_filled and weld_insp_filled and irn_filled:
            status = "Ready to Release"
        elif sdf["fitup_date"].eq("").all():
            status = "Not Started"
        else:
            status = "Under Fabrication"

        status_map[key] = status

    df["Spool Status"] = df["spool_key"].map(status_map)
    df["joint_size"] = pd.to_numeric(df["joint_size"], errors="coerce")
    return df


_CLASSIFIED_COLS = [
    "spool_key", "wo_no", "batch_no", "material_group", "zone", "service", "line_no",
    "iso_dwg_no", "dwg_spool_no", "iso_run_no", "rev", "shop_field", "joint_no",
    "joint_size", "welding_type", "pwht", "fitup_date", "welding_date",
    "fitup_inspection_date", "welding_inspection_date", "irn_date",
    "delivery_date", "site_delivery_date", "paint_system", "paint_status",
    "Spool Status",
]

# progression order used for the Summary sheet
_STATUS_ORDER = [
    "Not Started", "Under Fabrication", "Ready for PWHT", "All done-Awaiting IRN",
    "Ready to Release", "Ready to Release-Straight Pipe",
    "Sent to Painting", "Sent to Site",
]

_FILL_COLORS = {
    "Ready to Release": "C6EFCE",
    "Ready to Release-Straight Pipe": "D5F5E3",
    "Under Fabrication": "FFEB9C",
    "Sent to Painting": "F4B084",
    "Sent to Site": "D9D2E9",
    "Not Started": "FFC7CE",
    "Ready for PWHT": "A9D0F5",
    "All done-Awaiting IRN": "FADBD8",
}


def _format_sheet(ws) -> None:
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    thin = Side(border_style="thin")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        status = row[-1].value if len(row) > 1 else None
        code = _FILL_COLORS.get(status, "FFFFFF")
        fill = PatternFill(start_color=code, end_color=code, fill_type="solid")
        for cell in row:
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in ws.columns:
        width = max(len(str(c.value) or "") for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = width


def summarize(classified: pd.DataFrame) -> pd.DataFrame:
    shop = classified[
        (classified["shop_field"] == "S")
        | classified["Spool Status"].isin(["Ready to Release-Straight Pipe", "Sent to Site"])
    ]
    g = (
        shop.groupby("Spool Status")
        .agg(**{"Pipe Spools": ("spool_key", "nunique"),
                "Joints": ("joint_no", "count"),
                "Dia-Inch": ("joint_size", "sum")})
        .reset_index()
    )
    tot_s, tot_d = g["Pipe Spools"].sum(), g["Dia-Inch"].sum()
    g["% Spools"] = (g["Pipe Spools"] / tot_s * 100).round(1) if tot_s else 0.0
    g["% Dia-Inch"] = (g["Dia-Inch"] / tot_d * 100).round(1) if tot_d else 0.0
    g["Dia-Inch"] = g["Dia-Inch"].round(2)
    _rank = {s: i for i, s in enumerate(_STATUS_ORDER)}
    return (g.sort_values("Spool Status", key=lambda c: c.map(lambda x: _rank.get(x, 99)))
             .reset_index(drop=True)
             [["Spool Status", "Pipe Spools", "Joints", "Dia-Inch", "% Spools", "% Dia-Inch"]])


def build_classified_xlsx(df: pd.DataFrame) -> bytes:
    classified = classify(df)
    df_final = classified[[c for c in _CLASSIFIED_COLS if c in classified.columns]].copy()
    df_shop = df_final[
        (df_final["shop_field"] == "S")
        | df_final["Spool Status"].isin(["Ready to Release-Straight Pipe", "Sent to Site"])
    ]
    status_summary = summarize(classified)
    shop_field_summary = (
        classified.groupby("shop_field").agg(Total_DiaInch=("joint_size", "sum")).reset_index()
    )
    _pss_cols = ["spool_key", "wo_no", "batch_no", "zone", "material_group", "iso_dwg_no",
                 "line_no", "dwg_spool_no", "iso_run_no", "paint_system", "paint_status",
                 "Spool Status"]
    pipe_spool_summary = df_shop.drop_duplicates(subset=["spool_key"])[
        [c for c in _pss_cols if c in df_shop.columns]
    ]

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df_final.to_excel(writer, index=False, sheet_name="All Spools")
        pipe_spool_summary.to_excel(writer, index=False, sheet_name="Pipe Spool Summary")
        shop_field_summary.to_excel(writer, index=False, sheet_name="Shop and Field")
        for status in _STATUS_ORDER:
            part = df_shop[df_shop["Spool Status"] == status]
            if not part.empty:
                part.to_excel(writer, index=False, sheet_name=status[:31])
        wb = writer.book
        _write_summary_sheet(wb, status_summary)      # proper, formatted, tab 1
        for name in wb.sheetnames:
            if name != "Summary":
                _format_sheet(wb[name])
    return buf.getvalue()


def _write_summary_sheet(wb, summ: pd.DataFrame) -> None:
    """Formatted Summary tab: title, colour-keyed status rows, TOTAL row."""
    ws = wb.create_sheet("Summary", 0)
    thin = Side(border_style="thin", color="BFBFBF")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    headers = list(summ.columns)
    HR = 4                                            # header row

    ws["A1"] = "CLASSIFIED SPOOLS  —  SUMMARY"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A2"] = (f"Generated {datetime.now(MYT):%Y-%m-%d %H:%M} MYT   ·   "
                "shop spools plus straight pipe / sent-to-site")
    ws["A2"].font = Font(italic=True, size=9, color="808080")

    for j, h in enumerate(headers, 1):
        c = ws.cell(HR, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = box

    r = HR + 1
    for _, row in summ.iterrows():
        for j, h in enumerate(headers, 1):
            c = ws.cell(r, j, row[h])
            c.border = box
            c.alignment = Alignment(horizontal="left" if h == "Spool Status" else "center")
            if h == "Dia-Inch":
                c.number_format = "#,##0.00"
            elif h.startswith("%"):
                c.number_format = '0.0"%"'
            elif h in ("Pipe Spools", "Joints"):
                c.number_format = "#,##0"
        code = _FILL_COLORS.get(row["Spool Status"], "FFFFFF")
        ws.cell(r, 1).fill = PatternFill("solid", start_color=code)
        r += 1

    ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    ws.cell(r, 1).border = box
    for j, h in enumerate(headers[1:], 2):
        col = ws.cell(HR, j).column_letter
        cell = ws.cell(r, j)
        if h.startswith("%"):
            cell.value = 100.0
            cell.number_format = '0.0"%"'
        else:
            cell.value = f"=SUM({col}{HR + 1}:{col}{r - 1})"
            cell.number_format = "#,##0.00" if h == "Dia-Inch" else "#,##0"
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = box

    for j, h in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(HR, j).column_letter].width = \
            max(len(h) + 4, 22 if h == "Spool Status" else 13)
    ws.freeze_panes = f"A{HR + 1}"


# ----------------------------------------------------------------------
# master export  (mirrors export_master.py)
# ----------------------------------------------------------------------
_RENAME = {
    "wo_no": "WO NO", "zone": "ZONE", "status": "STATUS", "workable": "WORKABLE",
    "batch_no": "BATCH NO.", "material_group": "MATERIAL GROUP", "area": "AREA",
    "location": "LOCATION", "service": "SERVICE", "line_no": "LINE NO.",
    "line_spec": "LINE SPEC.", "iso_dwg_no": "ISO DWG NO.", "dwg_spool_no": "DWG SPOOL NO",
    "iso_run_no": "ISO RUN NO.", "rev": "REV", "shop_field": "SHOP/FIELD",
    "joint_no": "JOINT NO", "item_1": "ITEM 1", "sch_rating_1": "SCH / RATING 1",
    "heat_no_1": "HEAT NO. 1", "item_2": "ITEM 2", "sch_rating_2": "SCH / RATING 2",
    "heat_no_2": "HEAT NO. 2", "joint_size": "JOINT SIZE", "welding_type": "WELDING TYPE",
    "paint_system": "PAINT SYSTEM", "paint_status": "PAINT STATUS",
    "pwht": "PWHT", "fitup_date": "FIT UP DATE",
    "fitup_inspection_date": "FIT UP INSPECTION DATE", "fu_report_no": "FU REPORT NO",
    "welding_date": "WELDING DATE", "welding_inspection_date": "WELDING INSPECTION DATE",
    "root_welder_no": "ROOT WELDER NO", "capping_welder_no": "CAPPING WELDER NO",
    "visual_report_no": "VISUAL REPORT NO", "welding_process": "WELDING PROCESS",
    "wps_no": "WPS NO", "rt_bsr_date": "RT BSR DATE", "rt_bsr_report_no": "RT BSR REPORT NO",
    "pwht_date": "PWHT DATE", "pwht_report_no": "PWHT REPORT NO", "rt_asr_date": "RT ASR DATE",
    "rt_asr_report_no": "RT ASR REPORT NO", "mpi_pt_date": "MPI PT DATE",
    "mpi_pt_report_no": "MPI PT REPORT NO", "hardness_date": "HARDNESS DATE",
    "hardness_report_no": "HARNESS REPORT NO", "pmi_date": "PMI DATE",
    "pmi_report_no": "PMI REPORT NO", "ferrite_date": "FERRITE DATE",
    "ferrite_report_no": "FERRITE REPORT NO", "irn_date": "IRN DATE",
    "irn_report_no": "IRN REPORT NO", "delivery_order_no": "PAINTING DO NO",
    "delivery_date": "PAINTING DELIVERY DATE", "site_do_no": "SITE DO NO",
    "site_delivery_date": "SITE DELIVERY DATE", "test_pack_no": "TEST PACK NO",
    "system_no": "SYSTEM NO", "subsystem_no": "SUB SYSTEM NO",
    "bsr_fresh_joint_status": "BSR FRESH JOINT STATUS",
    "bsr_repair_one_status": "BSR REPAIR ONE STATUS",
    "bsr_repair_two_status": "BSR REPAIR TWO STATUS", "schedule": "SCH",
    "rt_asr_result": "RT ASR RESULT", "mpi_pt_result": "MPI PT RESULT",
    "hardness_result": "HARDNESS RESULT", "pmi_result": "PMI RESULT", "remark": "REMARK",
    "test_pressure": "TEST PRESSURE", "total_film": "TOTAL FILM", "film_acc": "FILM ACC",
    "film_rej": "FILM REJ", "length_rej": "LENGTH REJ", "mpi_pt_type": "MPI PT TYPE",
    "painting_date": "PAINTING DATE", "painting_report_no": "PAINTING REPORT NO",
}

_ORDER = [
    "WO NO", "ZONE", "STATUS", "WORKABLE", "BATCH NO.", "AREA", "LOCATION", "SERVICE",
    "ISO DWG NO.", "ISO RUN NO.", "REV", "TEST PACK NO", "SYSTEM NO", "SUB SYSTEM NO",
    "TEST PRESSURE", "LINE NO.", "LINE SPEC.", "DWG SPOOL NO", "MATERIAL GROUP",
    "SHOP/FIELD", "JOINT NO", "JOINT SIZE", "SCH", "WPS NO", "WELDING PROCESS",
    "WELDING TYPE", "FIT UP INSPECTION DATE", "ITEM 1", "HEAT NO. 1", "ITEM 2",
    "HEAT NO. 2", "FU REPORT NO", "WELDING INSPECTION DATE", "ROOT WELDER NO",
    "CAPPING WELDER NO", "VISUAL REPORT NO", "RT BSR DATE", "RT BSR REPORT NO",
    "BSR FRESH JOINT STATUS", "TOTAL FILM", "FILM ACC", "FILM REJ", "LENGTH REJ",
    "BSR REPAIR ONE STATUS", "BSR REPAIR TWO STATUS", "PWHT DATE", "PWHT REPORT NO",
    "RT ASR DATE", "RT ASR REPORT NO", "RT ASR RESULT", "MPI PT DATE", "MPI PT TYPE",
    "MPI PT REPORT NO", "MPI PT RESULT", "HARDNESS DATE", "HARNESS REPORT NO",
    "HARDNESS RESULT", "PMI DATE", "PMI REPORT NO", "PMI RESULT", "FERRITE DATE",
    "FERRITE REPORT NO", "PAINTING DATE", "PAINTING REPORT NO", "IRN DATE",
    "IRN REPORT NO", "PAINT SYSTEM", "PAINT STATUS", "PWHT",
    "SCH / RATING 1", "SCH / RATING 2",
    "FIT UP DATE", "WELDING DATE", "PAINTING DO NO", "PAINTING DELIVERY DATE",
    "SITE DO NO", "SITE DELIVERY DATE", "REMARK",
]


def _strip_tz(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for c in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[c]):
            s = df[c]
            if getattr(s.dtype, "tz", None) is not None:
                s = s.dt.tz_localize(None)
            df[c] = s
    return df


def build_master_xlsx(df: pd.DataFrame) -> bytes:
    df = _strip_tz(df)
    df = df.rename(columns={k: v for k, v in _RENAME.items() if k in df.columns})
    df = df.replace("nan", "").fillna("")
    remaining = [c for c in df.columns if c not in _ORDER]
    df = df.reindex(columns=[c for c in _ORDER if c in df.columns] + remaining)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Master Data")
        ws = writer.sheets["Master Data"]
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max((len(str(c.value)) if c.value else 0) for c in col) + 2
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
    return buf.getvalue()


def stamp() -> str:
    """Timestamp for file names, in Malaysia time (server clock is UTC)."""
    return datetime.now(MYT).strftime("%Y-%m-%d_%H%M")


# ----------------------------------------------------------------------
# Excel -> spools import  (mirrors tabs/inventory_tab.py import_excel_to_db)
# Excel header -> db column
# ----------------------------------------------------------------------
IMPORT_MAP = {
    "WO NO": "wo_no", "ZONE": "zone", "STATUS": "status", "WORKABLE": "workable",
    "BATCH NO.": "batch_no", "AREA": "area", "LOCATION": "location", "SERVICE": "service",
    "ISO DWG NO.": "iso_dwg_no", "ISO RUN NO.": "iso_run_no", "REV": "rev",
    "TEST PACK NO": "test_pack_no", "SYSTEM NO": "system_no", "SUB SYSTEM NO": "subsystem_no",
    "TEST PRESSURE": "test_pressure", "LINE NO.": "line_no", "LINE SPEC.": "line_spec",
    "DWG SPOOL NO": "dwg_spool_no", "MATERIAL GROUP": "material_group",
    "SHOP/FIELD": "shop_field", "JOINT NO": "joint_no", "JOINT SIZE": "joint_size",
    "SCH": "schedule", "WPS NO": "wps_no", "WELDING PROCESS": "welding_process",
    "WELDING TYPE": "welding_type", "FIT UP INSPECTION DATE": "fitup_inspection_date",
    "ITEM 1": "item_1", "SCH / RATING 1": "sch_rating_1", "HEAT NO. 1": "heat_no_1",
    "ITEM 2": "item_2", "SCH / RATING 2": "sch_rating_2", "HEAT NO. 2": "heat_no_2",
    "FU REPORT NO": "fu_report_no", "WELDING INSPECTION DATE": "welding_inspection_date",
    "ROOT WELDER NO": "root_welder_no", "CAPPING WELDER NO": "capping_welder_no",
    "VISUAL REPORT NO": "visual_report_no", "RT BSR DATE": "rt_bsr_date",
    "RT BSR REPORT NO": "rt_bsr_report_no", "BSR FRESH JOINT STATUS": "bsr_fresh_joint_status",
    "TOTAL FILM": "total_film", "FILM ACC": "film_acc", "FILM REJ": "film_rej",
    "LENGTH REJ": "length_rej", "BSR REPAIR ONE STATUS": "bsr_repair_one_status",
    "BSR REPAIR TWO STATUS": "bsr_repair_two_status", "PWHT DATE": "pwht_date",
    "PWHT REPORT NO": "pwht_report_no", "RT ASR DATE": "rt_asr_date",
    "RT ASR REPORT NO": "rt_asr_report_no", "RT ASR RESULT": "rt_asr_result",
    "MPI PT DATE": "mpi_pt_date", "MPI PT TYPE": "mpi_pt_type",
    "MPI PT REPORT NO": "mpi_pt_report_no", "MPI PT RESULT": "mpi_pt_result",
    "HARDNESS DATE": "hardness_date", "HARNESS REPORT NO": "hardness_report_no",
    "HARDNESS RESULT": "hardness_result", "PMI DATE": "pmi_date",
    "PMI REPORT NO": "pmi_report_no", "PMI RESULT": "pmi_result",
    "FERRITE DATE": "ferrite_date", "FERRITE REPORT NO": "ferrite_report_no",
    "PAINTING DATE": "painting_date", "PAINTING REPORT NO": "painting_report_no",
    "IRN DATE": "irn_date", "IRN REPORT NO": "irn_report_no", "PAINT SYSTEM": "paint_system",
    "PAINT STATUS": "paint_status",
    "PWHT": "pwht", "FIT UP DATE": "fitup_date", "WELDING DATE": "welding_date",
    "PAINTING DO NO": "delivery_order_no", "PAINTING DELIVERY DATE": "delivery_date",
    "SITE DO NO": "site_do_no", "SITE DELIVERY DATE": "site_delivery_date",
    "REMARK": "remark",
}


def spools_df_from_excel(raw: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Map an uploaded master-format sheet to spools columns.
    Returns (clean_df, missing_headers)."""
    missing = [h for h in IMPORT_MAP if h not in raw.columns]
    df = raw.rename(columns=IMPORT_MAP)
    cols = [c for c in dict.fromkeys(IMPORT_MAP.values()) if c in df.columns]
    df = df[cols].copy()
    if "joint_size" in df.columns:
        df["joint_size"] = pd.to_numeric(df["joint_size"], errors="coerce")

    _NULLISH = {"", "nan", "nat", "none", "null", "#n/a", "n/a"}
    text_cols = [c for c in cols if c != "joint_size"]
    # object-first so NaN/NaT reliably become None (not the string "nan")
    df[text_cols] = df[text_cols].astype(object).where(pd.notna(df[text_cols]), None)
    for c in text_cols:
        df[c] = df[c].map(
            lambda v: None if v is None or str(v).strip().lower() in _NULLISH
            else str(v).strip()
        )
    return df, missing


def summary_status_counts(df: pd.DataFrame) -> dict[str, int]:
    """Spool-status tally, mirroring tabs/project_summary_tab.py (5 buckets)."""
    d = df[df["shop_field"] == "S"].copy()
    for c in ("line_no", "iso_dwg_no", "dwg_spool_no", "iso_run_no"):
        d[c] = d[c].fillna("").astype(str).str.strip()
    d["k"] = d["line_no"] + "_" + d["iso_dwg_no"] + "_" + d["dwg_spool_no"] + "_" + d["iso_run_no"]

    counter = {"Not Started": 0, "Under Fabrication": 0, "Ready to Release": 0,
               "Sent to Painting": 0, "Sent to Site": 0}
    ss = {"SS", "SS304", "SS316"}

    def filled(s):
        return s.fillna("").astype(str).str.strip().ne("")

    for _, g in d.groupby("k"):
        mat = str(g["material_group"].fillna("").iloc[0]).upper()
        spool = str(g["dwg_spool_no"].iloc[0])
        all_f = lambda c: bool(filled(g[c]).all())
        any_f = lambda c: bool(filled(g[c]).any())

        if spool.startswith("SP-SPL"):
            s = "Ready to Release"
        elif any_f("site_delivery_date"):
            s = "Sent to Site"
        elif any_f("delivery_date"):
            s = "Sent to Site" if mat in ss else "Sent to Painting"
        elif all_f("fitup_inspection_date") and all_f("welding_inspection_date") and all_f("irn_date"):
            s = "Ready to Release"
        elif not any_f("fitup_date"):
            s = "Not Started"
        else:
            s = "Under Fabrication"
        counter[s] += 1
    return counter


def build_full_backup_xlsx(tables: dict[str, pd.DataFrame]) -> bytes:
    """One workbook, one sheet per table - a portable snapshot."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, tdf in tables.items():
            _strip_tz(tdf).to_excel(writer, index=False, sheet_name=name[:31])
    return buf.getvalue()
