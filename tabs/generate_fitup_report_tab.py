import tkinter as tk
from tkinter import ttk, messagebox
import sqlite3
import pandas as pd
import os
from openpyxl import load_workbook
from tkinter.filedialog import asksaveasfilename

def build_tab(parent, permissions):
    frame = ttk.Frame(parent)

    # ─── Entry Frame ─────────────────────────────────────────────────────────────
    entry_frame = ttk.LabelFrame(frame, text="Generate Fit-Up Report")
    entry_frame.pack(padx=20, pady=20, fill="x")

    ttk.Label(entry_frame, text="Fit-Up Report No:").grid(row=0, column=0, padx=5, pady=10, sticky="e")
    report_no_var = tk.StringVar()
    report_no_entry = ttk.Entry(entry_frame, textvariable=report_no_var, width=30)
    report_no_entry.grid(row=0, column=1, padx=5, pady=10)

    def generate_report():
        fu_report_no = report_no_var.get().strip()

        if not fu_report_no:
            messagebox.showerror("Missing Input", "Please enter Fit-Up Report No.")
            return

        try:
            conn = sqlite3.connect("database/spool_tracking.db")
            query = """
                SELECT 
                    iso_dwg_no,
                    iso_run_no,
                    line_no,
                    dwg_spool_no,
                    joint_no,
                    heat_no_1,
                    heat_no_2,
                    joint_size,
                    sch_rating_1,
                    welding_type
                FROM spools
                WHERE shop_field = 'S' AND fu_report_no = ?
            """
            df = pd.read_sql_query(query, conn, params=(fu_report_no,))
            conn.close()

            if df.empty:
                messagebox.showinfo("No Data", f"No records found for Fit-Up Report No: {fu_report_no}")
                return

            # Load Excel Template
            template_path = os.path.join("QC Report template", "Fit Up-Inspection Report Piping.xlsx")
            wb = load_workbook(template_path)
            ws = wb.active

            # Set Report No
            ws["M7"] = fu_report_no

            start_row = 11
            for i, row in df.iterrows():
                ws[f"A{start_row+i}"] = row['iso_dwg_no']
                ws[f"F{start_row+i}"] = row['iso_run_no']
                ws[f"G{start_row+i}"] = row['line_no']
                ws[f"I{start_row+i}"] = row['dwg_spool_no']
                ws[f"J{start_row+i}"] = row['joint_no']
                ws[f"K{start_row+i}"] = row['heat_no_1']
                ws[f"L{start_row+i}"] = row['heat_no_2']
                ws[f"M{start_row+i}"] = row['joint_size']
                ws[f"N{start_row+i}"] = row['sch_rating_1']
                ws[f"O{start_row+i}"] = "✓"
                ws[f"Q{start_row+i}"] = row['welding_type']
                # Pastikan A tidak wrap text
                ws[f"A{start_row+i}"].alignment = ws[f"A{start_row+i}"].alignment.copy(wrapText=False)

            # Simpan fail baru
            save_path = asksaveasfilename(defaultextension=".xlsx",
                                           filetypes=[("Excel Files", "*.xlsx")],
                                           initialfile=f"Fit-Up Report - {fu_report_no}.xlsx")
            if save_path:
                wb.save(save_path)
                messagebox.showinfo("Success", f"Report saved successfully:\n{save_path}")

        except Exception as e:
            messagebox.showerror("Error", f"Failed to generate report:\n{e}")

    ttk.Button(entry_frame, text="Generate Report", command=generate_report).grid(row=0, column=2, padx=10)

    return frame
