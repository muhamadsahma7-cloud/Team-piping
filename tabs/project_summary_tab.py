import tkinter as tk
from tkinter import ttk
import sqlite3
from datetime import datetime

def build_project_summary_tab(parent, permissions):
    container = tk.Frame(parent)

    canvas = tk.Canvas(container)
    scrollbar = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
    scrollable_frame = tk.Frame(canvas)

    scrollable_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
    canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
    canvas.configure(yscrollcommand=scrollbar.set)

    canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")
    container.pack(fill="both", expand=True)

    frame = scrollable_frame

    tk.Label(frame, text="Progress Summary", font=("Segoe UI", 14, "bold")).grid(row=0, column=0, padx=10, pady=10, sticky="w")

    tree1 = ttk.Treeview(frame, columns=("desc", "value"), show="headings", height=10)
    tree1.heading("desc", text="Description")
    tree1.heading("value", text="Value")
    tree1.grid(row=1, column=0, padx=10, sticky="nsew")

    tk.Label(frame, text="Work Order Issuance", font=("Segoe UI", 14, "bold")).grid(row=2, column=0, padx=10, pady=(20, 10), sticky="w")

    tree2 = ttk.Treeview(frame, columns=("wo_no", "material_group", "total", "fitup_balance", "weld_balance"), show="headings", height=10)
    tree2.heading("wo_no", text="WO No")
    tree2.heading("material_group", text="Material Group")
    tree2.heading("total", text="Total Dia Inch")
    tree2.heading("fitup_balance", text="Balance Fit-Up")
    tree2.heading("weld_balance", text="Balance Welding")
    tree2.grid(row=3, column=0, padx=10, sticky="nsew")

    tk.Label(frame, text="Work Order Summary", font=("Segoe UI", 14, "bold")).grid(row=4, column=0, padx=10, pady=(20, 10), sticky="w")

    tree3 = ttk.Treeview(frame, columns=("desc", "value"), show="headings", height=7)
    tree3.heading("desc", text="Description")
    tree3.heading("value", text="Value")
    tree3.grid(row=5, column=0, padx=10, sticky="nsew")

    def refresh_data():
        tree1.delete(*tree1.get_children())
        tree2.delete(*tree2.get_children())
        tree3.delete(*tree3.get_children())

        today = datetime.today().strftime("%Y-%m-%d")
        conn = sqlite3.connect("database/spool_tracking.db")
        cur = conn.cursor()

        def get_sum(query, params=()):
            cur.execute(query, params)
            result = cur.fetchone()
            return result[0] if result and result[0] else 0

        fitup_today = get_sum("SELECT SUM(joint_size) FROM spools WHERE shop_field = 'S' AND fitup_date = ?", (today,))
        weld_today = get_sum("SELECT SUM(joint_size) FROM spools WHERE shop_field = 'S' AND welding_date = ?", (today,))
        cum_fitup = get_sum("SELECT SUM(joint_size) FROM spools WHERE shop_field = 'S' AND TRIM(fitup_date) != ''")
        cum_welding = get_sum("SELECT SUM(joint_size) FROM spools WHERE shop_field = 'S' AND TRIM(welding_date) != ''")
        total_dia = get_sum("SELECT SUM(joint_size) FROM spools WHERE shop_field = 'S'")
        gap_fitup = cum_fitup - get_sum("SELECT SUM(joint_size) FROM spools WHERE shop_field = 'S' AND TRIM(fitup_inspection_date) != ''")
        gap_welding = cum_welding - get_sum("SELECT SUM(joint_size) FROM spools WHERE shop_field = 'S' AND TRIM(welding_inspection_date) != ''")

        cur.execute("SELECT * FROM spools WHERE shop_field = 'S'")
        rows = cur.fetchall()
        columns = [desc[0] for desc in cur.description]
        data = [dict(zip(columns, row)) for row in rows]

        spool_map = {}
        for row in data:
            key = f"{row['line_no']}_{row['iso_dwg_no']}_{row['dwg_spool_no']}_{row['iso_run_no']}"
            if key not in spool_map:
                spool_map[key] = []
            spool_map[key].append(row)

        status_counter = {
            "Not Started": 0,
            "Under Fabrication": 0,
            "Ready to Release": 0,
            "Sent to Painting": 0,
            "Sent to Site": 0
        }

        for key, joints in spool_map.items():
            material_group = (joints[0].get("material_group") or "").upper()
            dwg_spool_no = joints[0].get("dwg_spool_no") or ""

            def is_filled(field): return all(j.get(field) for j in joints)
            def any_filled(field): return any(j.get(field) for j in joints)

            if dwg_spool_no.startswith("SP-SPL"):
                status = "Ready to Release"
            elif any_filled("site_delivery_date"):
                status = "Sent to Site"
            elif any_filled("delivery_date"):
                status = "Sent to Site" if material_group in ["SS", "SS304", "SS316"] else "Sent to Painting"
            elif is_filled("fitup_inspection_date") and is_filled("welding_inspection_date") and is_filled("irn_date"):
                status = "Ready to Release"
            elif not any_filled("fitup_date"):
                status = "Not Started"
            elif not is_filled("welding_date"):
                status = "Under Fabrication"
            else:
                status = "Under Fabrication"

            status_counter[status] += 1

        summary_rows = [
            ("Today's Date", datetime.today().strftime("%d-%m-%Y (%A)")),
            ("Today's Fit-Up (Dia Inch)", fitup_today),
            ("Today's Welding (Dia Inch)", weld_today),
            ("Cumulative Fit-Up (Dia Inch)", cum_fitup),
            ("Cumulative Welding (Dia Inch)", cum_welding),
            ("Balance Fit-Up (Dia Inch)", total_dia - cum_fitup),
            ("Balance Welding (Dia Inch)", total_dia - cum_welding),
            ("Gap Fit-Up (Dia Inch)", gap_fitup),
            ("Gap Welding (Dia Inch)", gap_welding),
            ("Spool Not Started", status_counter["Not Started"]),
            ("Spool Under Fabrication", status_counter["Under Fabrication"]),
            ("Spool Ready to Release", status_counter["Ready to Release"]),
            ("Spool Sent to Painting", status_counter["Sent to Painting"]),
            ("Spool Sent to Site", status_counter["Sent to Site"]),
        ]
        for desc, val in summary_rows:
            tree1.insert("", "end", values=(desc, val))

        cur.execute("""
            SELECT wo_no, material_group, SUM(joint_size),
                   SUM(CASE WHEN fitup_date IS NULL OR TRIM(fitup_date) = '' THEN joint_size ELSE 0 END) as fitup_balance,
                   SUM(CASE WHEN welding_date IS NULL OR TRIM(welding_date) = '' THEN joint_size ELSE 0 END) as weld_balance
            FROM spools
            WHERE shop_field = 'S' AND LOWER(status) = 'issued'
            GROUP BY wo_no, material_group
            ORDER BY SUM(CASE WHEN welding_date IS NULL OR TRIM(welding_date) = '' THEN joint_size ELSE 0 END) DESC
        """)
        for wo_no, mat_grp, total, fitup_balance, weld_balance in cur.fetchall():
            row = tree2.insert("", "end", values=(wo_no, mat_grp, total, fitup_balance, weld_balance))
            if weld_balance == 0 and fitup_balance == 0:
                tree2.item(row, tags=("done",))
        tree2.tag_configure("done", background="#c8e6c9")

        tree3_data = [
            ("Total Dia Inch in Database", total_dia),
            ("Total Dia Inch Issued Work Order", get_sum("""
                SELECT SUM(joint_size) FROM spools
                WHERE shop_field = 'S' AND LOWER(status) = 'issued'
            """)),
            ("Balance Dia Inch Fit-Up (Issued W.O.)", get_sum("""
                SELECT SUM(joint_size) FROM spools
                WHERE shop_field = 'S' AND LOWER(status) = 'issued'
                AND (fitup_date IS NULL OR TRIM(fitup_date) = '')
            """)),
            ("Balance Dia Inch Welding (Issued W.O.)", get_sum("""
                SELECT SUM(joint_size) FROM spools
                WHERE shop_field = 'S' AND LOWER(status) = 'issued'
                AND (welding_date IS NULL OR TRIM(welding_date) = '')
            """)),
            ("Total Workable Dia Inch", get_sum("""
                SELECT SUM(joint_size) FROM spools
                WHERE shop_field = 'S' AND UPPER(TRIM(workable)) = 'Y'
            """)),
            ("Non Workable Dia Inch", get_sum("""
                SELECT SUM(joint_size) FROM spools
                WHERE shop_field = 'S' AND UPPER(TRIM(workable)) = 'N'
            """)),
            ("Total Dia Inch Unissued Work Order", get_sum("""
                SELECT SUM(joint_size) FROM spools
                WHERE shop_field = 'S' AND LOWER(status) = 'unissued'
            """)),
            ("Total Dia Inch Hold Work Order", get_sum("""
                SELECT SUM(joint_size) FROM spools
                WHERE shop_field = 'S' AND LOWER(status) = 'hold'
            """)),
            ("Total Dia Inch Outstanding Materials (Status: OS)", get_sum("""
                SELECT SUM(joint_size) FROM spools
                WHERE shop_field = 'S' AND LOWER(status) = 'os'
            """))
        ]

        for desc, val in tree3_data:
            tree3.insert("", "end", values=(desc, val))

        conn.close()

    refresh_data()

    def export_to_excel():
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Side, Font
        from tkinter import filedialog

        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Progress Summary"
        ws_list = [("Progress Summary", tree1), ("Work Order Issuance", tree2), ("Work Order Summary", tree3)]

        for name, tree in ws_list:
            if name != "Progress Summary":
                ws = wb.create_sheet(title=name)
            else:
                ws = ws1

            for row_index, child in enumerate(tree.get_children(), start=1):
                values = tree.item(child)['values']
                ws.append(values)

            thin = Side(border_style="thin")
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
                for cell in row:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

            if ws.max_row >= 1:
                for cell in ws[1]:
                    cell.font = Font(bold=True)

            for col in ws.columns:
                max_length = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = max_length + 2

        default_name = f"project_summary_{datetime.today().strftime('%Y-%m-%d')}.xlsx"
        file_path = filedialog.asksaveasfilename(
            initialfile=default_name,
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            title="Save Excel File"
        )
        if file_path:
            wb.save(file_path)

    # ─── Export & Refresh Button ───
    button_frame = tk.Frame(frame)
    button_frame.grid(row=6, column=0, pady=20)

    export_btn = ttk.Button(button_frame, text="📁 Export to Excel", command=export_to_excel)
    export_btn.pack(side="left", padx=5)

    refresh_btn = ttk.Button(button_frame, text="🔄 Refresh", command=refresh_data)
    refresh_btn.pack(side="left", padx=5)

    return container
