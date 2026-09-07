import sqlite3
import pandas as pd
from datetime import datetime
import os
import subprocess
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# 1. Prepare output folder and filename
os.makedirs("exports", exist_ok=True)
timestamp = datetime.now().strftime("%Y-%m-%d_%H%M")
filename = f"exports/spool_tracking_{timestamp}.xlsx"

# 2. Load the entire 'spools' table into a DataFrame
try:
    conn = sqlite3.connect("database/spool_tracking.db")
    df = pd.read_sql_query("SELECT * FROM spools", conn)
    conn.close()
    print(f"Loaded {len(df)} records from database")
except Exception as e:
    print(f"ERROR: Could not connect to database: {e}")
    print("Please ensure the database file exists at: database/spool_tracking.db")
    exit(1)

# 3. Map DB column names → Excel headers
rename_map = {
    "wo_no":                   "WO NO",
    "zone":                    "ZONE",
    "status":                  "STATUS",
    "workable":                "WORKABLE",
    "batch_no":                "BATCH NO.",
    "material_group":          "MATERIAL GROUP",
    "area":                    "AREA",
    "location":                "LOCATION",
    "service":                 "SERVICE",
    "line_no":                 "LINE NO.",
    "line_spec":               "LINE SPEC.",
    "iso_dwg_no":              "ISO DWG NO.",
    "dwg_spool_no":            "DWG SPOOL NO",
    "iso_run_no":              "ISO RUN NO.",
    "rev":                     "REV",
    "shop_field":              "SHOP/FIELD",
    "joint_no":                "JOINT NO",
    "item_1":                  "ITEM 1",
    "sch_rating_1":            "SCH / RATING 1",
    "heat_no_1":               "HEAT NO. 1",
    "item_2":                  "ITEM 2",
    "sch_rating_2":            "SCH / RATING 2",
    "heat_no_2":               "HEAT NO. 2",
    "joint_size":              "JOINT SIZE",
    "welding_type":            "WELDING TYPE",
    "paint_system":            "PAINT SYSTEM",
    "pwht":                    "PWHT",
    "fitup_date":              "FIT UP DATE",
    "fitup_inspection_date":   "FIT UP INSPECTION DATE",
    "fu_report_no":            "FU REPORT NO",
    "welding_date":            "WELDING DATE",
    "welding_inspection_date": "WELDING INSPECTION DATE",
    "root_welder_no":          "ROOT WELDER NO",
    "capping_welder_no":       "CAPPING WELDER NO",
    "visual_report_no":        "VISUAL REPORT NO",
    "welding_process":         "WELDING PROCESS",
    "wps_no":                  "WPS NO",
    "rt_bsr_date":             "RT BSR DATE",
    "rt_bsr_report_no":        "RT BSR REPORT NO",
    "pwht_date":               "PWHT DATE",
    "pwht_report_no":          "PWHT REPORT NO",
    "rt_asr_date":             "RT ASR DATE",
    "rt_asr_report_no":        "RT ASR REPORT NO",
    "mpi_pt_date":             "MPI PT DATE",
    "mpi_pt_report_no":        "MPI PT REPORT NO",
    "hardness_date":           "HARDNESS DATE",
    "hardness_report_no":      "HARNESS REPORT NO",
    "pmi_date":                "PMI DATE",
    "pmi_report_no":           "PMI REPORT NO",
    "ferrite_date":            "FERRITE DATE",
    "ferrite_report_no":       "FERRITE REPORT NO",  
    "irn_date":                "IRN DATE",
    "irn_report_no":           "IRN REPORT NO",
    "delivery_order_no":       "PAINTING DO NO",
    "delivery_date":           "PAINTING DELIVERY DATE",
    "site_do_no":              "SITE DO NO",
    "site_delivery_date":      "SITE DELIVERY DATE",
    "test_pack_no":            "TEST PACK NO",
    "system_no":               "SYSTEM NO",
    "subsystem_no":            "SUB SYSTEM NO",
    "bsr_fresh_joint_status":  "BSR FRESH JOINT STATUS",
    "bsr_repair_one_status":   "BSR REPAIR ONE STATUS",
    "bsr_repair_two_status":   "BSR REPAIR TWO STATUS",
    "schedule":                "SCH",
    "rt_asr_result":           "RT ASR RESULT",
    "mpi_pt_result":           "MPI PT RESULT",
    "hardness_result":         "HARDNESS RESULT",
    "pmi_result":              "PMI RESULT",
    "remark":                  "REMARK",
    "test_pressure":           "TEST PRESSURE",
    "total_film":              "TOTAL FILM",
    "film_acc":                "FILM ACC",
    "film_rej":                "FILM REJ",
    "length_rej":              "LENGTH REJ",
    "mpi_pt_type":             "MPI PT TYPE",
    "painting_date":           "PAINTING DATE",
    "painting_report_no":      "PAINTING REPORT NO"
}

# 4. Only rename if the column actually exists in df
existing_rename = {orig: new for orig, new in rename_map.items() if orig in df.columns}
df = df.rename(columns=existing_rename)

# 5. Replace literal "nan" or NaN with empty string
df = df.replace("nan", "").fillna("")

# 6. Build final column order
excel_order = [
    "WO NO", "ZONE", "STATUS", "WORKABLE", "BATCH NO.", "AREA", "LOCATION", "SERVICE", "ISO DWG NO.", "ISO RUN NO.", "REV",
    "TEST PACK NO", "SYSTEM NO", "SUB SYSTEM NO", "TEST PRESSURE", "LINE NO.", "LINE SPEC.", "DWG SPOOL NO", "MATERIAL GROUP", "SHOP/FIELD", "JOINT NO",
    "JOINT SIZE", "SCH", "WPS NO", "WELDING PROCESS", "WELDING TYPE", "FIT UP INSPECTION DATE", "ITEM 1", "HEAT NO. 1",
    "ITEM 2", "HEAT NO. 2", "FU REPORT NO", "WELDING INSPECTION DATE", "ROOT WELDER NO", "CAPPING WELDER NO", "VISUAL REPORT NO", 
    "RT BSR DATE", "RT BSR REPORT NO", "BSR FRESH JOINT STATUS", "TOTAL FILM", "FILM ACC", "FILM REJ", "LENGTH REJ", "BSR REPAIR ONE STATUS", "BSR REPAIR TWO STATUS",
    "PWHT DATE", "PWHT REPORT NO", "RT ASR DATE", "RT ASR REPORT NO", "RT ASR RESULT",
    "MPI PT DATE", "MPI PT TYPE", "MPI PT REPORT NO", "MPI PT RESULT", "HARDNESS DATE", "HARNESS REPORT NO", "HARDNESS RESULT",
    "PMI DATE", "PMI REPORT NO", "PMI RESULT", "FERRITE DATE", "FERRITE REPORT NO", "PAINTING DATE", "PAINTING REPORT NO", "IRN DATE", "IRN REPORT NO",
    "PAINT SYSTEM", "PWHT", "SCH / RATING 1", "SCH / RATING 2", "FIT UP DATE", "WELDING DATE", "PAINTING DO NO", "PAINTING DELIVERY DATE", "SITE DO NO", "SITE DELIVERY DATE", "REMARK"
] 

remaining = [col for col in df.columns if col not in excel_order]
final_order = [col for col in excel_order if col in df.columns] + remaining
df = df.reindex(columns=final_order)

# 7. Write to Excel with formatting
with pd.ExcelWriter(filename, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Master Data")
    wb = writer.book
    ws = writer.sheets["Master Data"]

    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    for col in ws.columns:
        max_length = max((len(str(cell.value)) if cell.value else 0) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max_length + 2

# 8. Verify file was created successfully
if os.path.exists(filename):
    file_size = os.path.getsize(filename)
    print(f"Export complete: {filename}")
    print(f"File size: {file_size:,} bytes")
    
    # Auto-open the file
    try:
        os.startfile(os.path.abspath(filename))
        print("File opened successfully!")
    except Exception as e:
        print(f"Could not open the file automatically: {e}")
        print(f"Please manually open: {os.path.abspath(filename)}")
else:
    print(f"ERROR: Export failed - file was not created: {filename}")
