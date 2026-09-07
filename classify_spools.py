import sqlite3
import pandas as pd
import os
import time
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# ─── Setup ───
os.makedirs("classified_spools", exist_ok=True)
DB_PATH = r"P:\A_NAEC PROJECT FOLDER\NAEC-606_MIE-PREFCHEM UG PIPING\Team piping software\database\spool_tracking.db"
conn = sqlite3.connect(DB_PATH)
df = pd.read_sql_query("SELECT * FROM spools", conn)
conn.close()

# ─── Buat Spool Key ───
df["spool_key"] = (
    df["line_no"].astype(str).str.strip() + "_" +
    df["iso_dwg_no"].astype(str).str.strip() + "_" +
    df["dwg_spool_no"].astype(str).str.strip() + "_" +
    df["iso_run_no"].astype(str).str.strip()
)

# ─── Klasifikasi Status ───
status_map = {}
for spool in df["spool_key"].unique():
    spool_df = df[df["spool_key"] == spool]

    is_straight_pipe = spool_df["dwg_spool_no"].astype(str).str.startswith("SP-SPL").all()
    if not is_straight_pipe:
        spool_df = spool_df[spool_df["shop_field"] == "S"]
        if spool_df.empty:
            continue

    material_group = spool_df["material_group"].astype(str).str.upper().iloc[0]
    pwht_flag = spool_df["pwht"].astype(str).str.upper().iloc[0]
    site_date = spool_df.get("site_delivery_date")
    painting_date = spool_df.get("delivery_date")

    # Semua tarikh penuh
    fitup_filled = spool_df["fitup_date"].notna().all() and spool_df["fitup_date"].ne("").all()
    weld_filled = spool_df["welding_date"].notna().all() and spool_df["welding_date"].ne("").all()
    pwht_filled = "pwht_date" in spool_df.columns and spool_df["pwht_date"].notna().all() and spool_df["pwht_date"].ne("").all()
    fitup_insp_filled = spool_df["fitup_inspection_date"].notna().all() and spool_df["fitup_inspection_date"].ne("").all()
    weld_insp_filled = spool_df["welding_inspection_date"].notna().all() and spool_df["welding_inspection_date"].ne("").all()
    irn_filled = spool_df["irn_date"].notna().all() and spool_df["irn_date"].ne("").all()

    # ─── Logik Status ───
    if is_straight_pipe:
        if site_date.notna().any() and site_date.astype(str).str.strip().ne("").any():
            status = "Sent to Site"
        elif painting_date.notna().any() and painting_date.astype(str).str.strip().ne("").any():
            status = "Sent to Painting"
        else:
            status = "Ready to Release-Straight Pipe"
    elif site_date.notna().any() and site_date.astype(str).str.strip().ne("").any():
        status = "Sent to Site"
    elif painting_date.notna().any() and painting_date.astype(str).str.strip().ne("").any():
        if material_group in ["SS", "SS304", "SS316"]:
            status = "Sent to Site"
        else:
            status = "Sent to Painting"
    elif material_group in ["SS304", "SS_304", "SS316", "SS_316"] and fitup_insp_filled and weld_insp_filled and irn_filled:
        status = "Ready to Release"
    elif pwht_flag == "YES" and fitup_filled and weld_filled and not pwht_filled:
        status = "Ready for PWHT"
    elif (pwht_flag == "YES" and fitup_filled and weld_filled and pwht_filled) or \
         (pwht_flag != "YES" and fitup_filled and weld_filled):
        if not irn_filled:
            status = "All done-Awaiting IRN"
        else:
            status = "Ready to Release"
    elif fitup_insp_filled and weld_insp_filled and irn_filled:
        status = "Ready to Release"
    elif spool_df["fitup_date"].isnull().all() or spool_df["fitup_date"].eq("").all():
        status = "Not Started"
    elif spool_df["welding_date"].isnull().any() or spool_df["welding_date"].eq("").any():
        status = "Under Fabrication"
    else:
        status = "Under Fabrication"

    status_map[spool] = status

df["Spool Status"] = df["spool_key"].map(status_map)

# ─── Pilih Kolum untuk All Spools ───
df["joint_size"] = pd.to_numeric(df["joint_size"], errors="coerce")
df_final = df[[ 
    "spool_key", "wo_no", "material_group", "zone", "service", "line_no", "iso_dwg_no",
    "dwg_spool_no", "iso_run_no", "rev", "shop_field", "joint_no",
    "joint_size", "welding_type", "pwht", "fitup_date", "welding_date",
    "fitup_inspection_date", "welding_inspection_date", "irn_date",
    "delivery_date", "site_delivery_date", "paint_system", "Spool Status"
]].copy()

# ─── Penapisan untuk Summary ───
df_shop = df_final[
    (df_final["shop_field"] == "S") |
    df_final["Spool Status"].isin(["Ready to Release-Straight Pipe", "Sent to Site"])
]

# ─── Summary Status ───
status_summary = df_shop.groupby("Spool Status").agg(
    Total_Spools=('spool_key', 'nunique'),
    Total_Joints=('joint_no', 'count'),
    Total_DiaInch=('joint_size', 'sum')
).reset_index()

# ─── Shop and Field Summary ───
shop_field_summary = df.groupby("shop_field").agg(
    Total_DiaInch=('joint_size', 'sum')
).reset_index()

# ─── Pipe Spool Summary ───
pipe_spool_summary = df_shop.drop_duplicates(subset=["spool_key"]).copy()
pipe_spool_summary = pipe_spool_summary[[ 
    "spool_key", "wo_no", "zone", "material_group", "iso_dwg_no", "line_no", 
    "dwg_spool_no", "iso_run_no", "paint_system", "Spool Status"
]]

# ─── Export Excel ───
today_str = datetime.today().strftime("%Y-%m-%d_%H%M")
path = f"classified_spools/classified_spools_{today_str}.xlsx"
with pd.ExcelWriter(path, engine="openpyxl") as writer:
    status_summary.to_excel(writer, index=False, sheet_name="Summary")
    df_final.to_excel(writer, index=False, sheet_name="All Spools")
    pipe_spool_summary.to_excel(writer, index=False, sheet_name="Pipe Spool Summary")
    shop_field_summary.to_excel(writer, index=False, sheet_name="Shop and Field")
    for status in df_shop["Spool Status"].unique():
        df_status = df_shop[df_shop["Spool Status"] == status]
        df_status.to_excel(writer, index=False, sheet_name=status[:31])

# ─── Format Excel ───
wb = load_workbook(path)
fill_colors = {
    "Ready to Release": "C6EFCE",
    "Ready to Release-Straight Pipe": "D5F5E3",
    "Under Fabrication": "FFEB9C",
    "Sent to Painting": "F4B084",
    "Sent to Site": "D9D2E9",
    "Not Started": "FFC7CE",
    "Ready for PWHT": "A9D0F5",
    "All done-Awaiting IRN": "FADBD8"
}

def format_sheet(ws):
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    thin = Side(border_style="thin")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        status = row[-1].value if len(row) > 1 else None
        fill_code = fill_colors.get(status, "FFFFFF")
        fill = PatternFill(start_color=fill_code, end_color=fill_code, fill_type="solid")
        for cell in row:
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in ws.columns:
        max_len = max(len(str(cell.value) or "") for cell in col) + 2
        ws.column_dimensions[col[0].column_letter].width = max_len

for sheet in wb.sheetnames:
    if sheet != "Summary":
        format_sheet(wb[sheet])

wb.save(path)
print(f"✅ Excel exported: {path}")
time.sleep(1)
os.startfile(os.path.abspath(path))
